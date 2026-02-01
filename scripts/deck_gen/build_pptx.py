#!/usr/bin/env python3
"""
Build updated PowerPoint deck by applying LLM-generated updates.

Usage:
    python build_pptx.py <previous_deck> <updates_json> <output_pptx>
"""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import json
import sys
import shutil
import os

def build_updated_deck(previous_deck_path, updates_path, output_path):
    """
    Create new deck by:
    1. Copying previous deck as starting point
    2. Applying updates from LLM
    3. Saving as new file
    """
    
    # Copy previous deck as starting point
    print(f"📋 Copying previous deck as base...")
    shutil.copy(previous_deck_path, output_path)
    
    # Load the copy
    prs = Presentation(output_path)
    
    # Load updates
    with open(updates_path) as f:
        updates = json.load(f)
    
    print(f"🔄 Applying updates to {len(updates['slides'])} slides...")
    
    # Apply updates to each slide
    for slide_update in updates['slides']:
        slide_idx = slide_update['slide_index']
        
        if slide_idx >= len(prs.slides):
            print(f"⚠️  Warning: Slide index {slide_idx} out of range (deck has {len(prs.slides)} slides)")
            continue
        
        slide = prs.slides[slide_idx]
        
        print(f"   Slide {slide_idx}: {len(slide_update['updates'])} updates")
        
        for update in slide_update['updates']:
            try:
                apply_update(slide, update)
            except Exception as e:
                print(f"   ⚠️  Failed to apply update: {e}")
    
    # Save updated deck
    print(f"💾 Saving updated deck...")
    prs.save(output_path)
    
    return output_path

def replace_text_preserve_format(text_frame, old_text, new_text):
    """
    Replace text within a text frame while preserving run-level formatting.
    
    python-pptx gotcha: setting shape.text or paragraph.text replaces ALL runs
    with a single default-formatted run, nuking bold/italic/color/size.
    This function finds the run containing old_text and does an in-place
    replacement, keeping every run's font properties intact.
    
    Falls back to first-run replacement if old_text spans multiple runs.
    """
    # Fast path: single-run paragraph (most common in decks)
    for paragraph in text_frame.paragraphs:
        full_text = paragraph.text
        if old_text not in full_text:
            continue
        
        # Try run-level replacement first (preserves formatting perfectly)
        for run in paragraph.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text)
                return True
        
        # old_text spans multiple runs — rebuild using first run's formatting
        if paragraph.runs:
            replaced = full_text.replace(old_text, new_text)
            # Keep first run with new text, clear the rest
            paragraph.runs[0].text = replaced
            for run in paragraph.runs[1:]:
                run.text = ""
            return True
    
    return False


def apply_update(slide, update):
    """Apply a single update to a slide."""
    
    update_type = update.get('type')
    
    if update_type == 'title':
        # Update slide title — preserve formatting via run-level replacement
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            old_text = update.get('old_text', '')
            new_text = update.get('new_text', '')
            
            if old_text:
                if not replace_text_preserve_format(slide.shapes.title.text_frame, old_text, new_text):
                    # old_text not found — set new_text on first run to preserve font
                    tf = slide.shapes.title.text_frame
                    if tf.paragraphs and tf.paragraphs[0].runs:
                        tf.paragraphs[0].runs[0].text = new_text
                        for run in tf.paragraphs[0].runs[1:]:
                            run.text = ""
                    else:
                        slide.shapes.title.text = new_text
            else:
                # No old_text anchor — write to first run if possible
                tf = slide.shapes.title.text_frame
                if tf.paragraphs and tf.paragraphs[0].runs:
                    tf.paragraphs[0].runs[0].text = new_text
                    for run in tf.paragraphs[0].runs[1:]:
                        run.text = ""
                else:
                    slide.shapes.title.text = new_text
    
    elif update_type == 'table':
        # Update table data
        table_index = update.get('table_index', 0)
        row_updates = update.get('row_updates', [])
        
        tables_found = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                if tables_found == table_index:
                    table = shape.table
                    
                    for row_update in row_updates:
                        row_idx = row_update['row_index']
                        col_idx = row_update['col_index']
                        new_text = str(row_update['new_text'])
                        
                        if row_idx < len(table.rows) and col_idx < len(table.columns):
                            cell = table.cell(row_idx, col_idx)
                            # Preserve cell formatting via run-level update
                            if cell.has_text_frame and cell.text_frame.paragraphs:
                                para = cell.text_frame.paragraphs[0]
                                if para.runs:
                                    para.runs[0].text = new_text
                                    for run in para.runs[1:]:
                                        run.text = ""
                                else:
                                    cell.text = new_text
                            else:
                                cell.text = new_text
                    
                    break
                tables_found += 1
    
    elif update_type == 'text_box':
        # Update specific text box by shape_id — preserve formatting
        shape_id = update.get('shape_id')
        old_text = update.get('old_text', '')
        new_text = update.get('new_text', '')
        
        for shape in slide.shapes:
            if shape.shape_id == shape_id and hasattr(shape, 'text_frame'):
                if old_text:
                    if not replace_text_preserve_format(shape.text_frame, old_text, new_text):
                        # Fallback: write to first run
                        tf = shape.text_frame
                        if tf.paragraphs and tf.paragraphs[0].runs:
                            tf.paragraphs[0].runs[0].text = new_text
                            for run in tf.paragraphs[0].runs[1:]:
                                run.text = ""
                        else:
                            shape.text = new_text
                else:
                    tf = shape.text_frame
                    if tf.paragraphs and tf.paragraphs[0].runs:
                        tf.paragraphs[0].runs[0].text = new_text
                        for run in tf.paragraphs[0].runs[1:]:
                            run.text = ""
                    else:
                        shape.text = new_text
                break
    
    elif update_type == 'text_replace':
        # Find and replace text anywhere on slide — preserve formatting
        old_text = update.get('old_text', '')
        new_text = update.get('new_text', '')
        
        if not old_text:
            return
        
        for shape in slide.shapes:
            if hasattr(shape, 'text_frame') and old_text in shape.text:
                replace_text_preserve_format(shape.text_frame, old_text, new_text)

def validate_output(output_path):
    """Quick validation that output is a valid PowerPoint."""
    try:
        prs = Presentation(output_path)
        return True, f"Valid PowerPoint with {len(prs.slides)} slides"
    except Exception as e:
        return False, f"Invalid PowerPoint: {e}"

if __name__ == '__main__':
    if len(sys.argv) != 4:
        print("Usage: python build_pptx.py <previous_deck> <updates_json> <output_pptx>")
        sys.exit(1)
    
    previous_deck = sys.argv[1]
    updates_path = sys.argv[2]
    output_path = sys.argv[3]
    
    if not os.path.exists(previous_deck):
        print(f"❌ Error: Previous deck not found: {previous_deck}")
        sys.exit(1)
    
    if not os.path.exists(updates_path):
        print(f"❌ Error: Updates file not found: {updates_path}")
        sys.exit(1)
    
    try:
        result = build_updated_deck(previous_deck, updates_path, output_path)
        
        # Validate
        valid, message = validate_output(result)
        if valid:
            print(f"✅ {message}")
            print(f"📁 Created: {result}")
        else:
            print(f"❌ {message}")
            sys.exit(1)
            
    except Exception as e:
        print(f"❌ Error building deck: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
