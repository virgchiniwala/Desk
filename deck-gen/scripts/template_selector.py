#!/usr/bin/env python3
"""Select and analyze the appropriate PowerPoint template.

Usage:
    python template_selector.py --type sprint_review --templates-path <dir> --output <analysis_json>
"""

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

TEMPLATE_MAP = {
    "sprint_review": "sprint_review.pptx",
    "exec_review": "exec_review.pptx",
}


def analyze_template(pptx_path: str) -> dict:
    """Analyze a PPTX template and extract shape metadata for each slide.

    Returns a dict with 'slides' list, each containing shape names, types,
    and current text.

    Raises:
        FileNotFoundError: If template file doesn't exist.
        ValueError: If file is not a valid PPTX.
    """
    path = Path(pptx_path)
    if not path.exists():
        raise FileNotFoundError(f"Template not found: {pptx_path}")

    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise ValueError(f"Invalid PPTX: {e}") from e

    slides = []
    for i, slide in enumerate(prs.slides):
        shapes = []
        for shape in slide.shapes:
            entry = {
                "name": shape.name,
                "shape_id": shape.shape_id,
                "type": (
                    shape.shape_type.name
                    if hasattr(shape.shape_type, "name")
                    else str(shape.shape_type)
                ),
            }

            if shape.has_text_frame:
                entry["text"] = shape.text_frame.text
                entry["has_text"] = True
            elif shape.has_table:
                table = shape.table
                entry["has_table"] = True
                entry["rows"] = len(table.rows)
                entry["cols"] = len(table.columns)
                entry["cells"] = [
                    {"row": r, "col": c, "text": table.cell(r, c).text}
                    for r in range(len(table.rows))
                    for c in range(len(table.columns))
                ]
            else:
                entry["has_text"] = False
                entry["has_table"] = False

            shapes.append(entry)

        slides.append({"index": i, "shapes": shapes})

    return {
        "template_path": str(pptx_path),
        "slide_count": len(prs.slides),
        "slides": slides,
    }


def select_template(template_type: str, templates_path: str = "templates") -> str:
    """Resolve a template type to its full file path.

    Args:
        template_type: One of the keys in TEMPLATE_MAP.
        templates_path: Directory containing template files.

    Returns:
        Absolute path to the selected template.

    Raises:
        ValueError: If template_type is not recognized.
        FileNotFoundError: If the template file doesn't exist.
    """
    if template_type not in TEMPLATE_MAP:
        raise ValueError(
            f"Unknown template type: {template_type}. Options: {list(TEMPLATE_MAP.keys())}"
        )

    filename = TEMPLATE_MAP[template_type]
    full_path = Path(templates_path) / filename

    if not full_path.exists():
        raise FileNotFoundError(f"Template file not found: {full_path}")

    return str(full_path.resolve())


def main():
    parser = argparse.ArgumentParser(description="Select and analyze a deck template")
    parser.add_argument(
        "--type", required=True, choices=list(TEMPLATE_MAP.keys()), help="Template type"
    )
    parser.add_argument(
        "--template-path",
        help="Direct path to a template deck (.pptx). Overrides --type mapping if provided.",
    )
    parser.add_argument(
        "--templates-path",
        default="templates",
        help="Directory with templates (default: templates/)",
    )
    parser.add_argument("--output", "-o", help="Write analysis JSON to file (default: stdout)")
    args = parser.parse_args()

    try:
        template_path = args.template_path if args.template_path else select_template(args.type, args.templates_path)
        analysis = analyze_template(template_path)

        output = json.dumps(analysis, indent=2)
        if args.output:
            Path(args.output).write_text(output)
            print(f"Analysis written to {args.output}", file=sys.stderr)
        else:
            print(output)

    except (ValueError, FileNotFoundError) as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
